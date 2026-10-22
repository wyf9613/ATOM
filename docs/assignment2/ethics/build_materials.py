"""Build the ATOM ethics report and offline HTML deck from local evidence.

Run from any directory with python-docx and reportlab installed. Never opens
hardware, modifies raw data, or uploads/submits files.
"""

from __future__ import annotations

import csv
import hashlib
import html
import json
import re
import statistics
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, PageBreak, KeepTogether


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
CONTENT = json.loads((HERE / "content.json").read_text())
ORAL = json.loads((HERE / "oral.json").read_text())
RAW = ROOT / "docs/data/gripper/20261007/empty_closure_20261007"
FONT_ROOT = Path("/System/Library/Fonts/Supplemental")
NAVY = "142C3E"
TEAL = "087F83"


def words(s):
    return len(s.split())


def conservative_words(s):
    # Splits hyphenated and slash-separated terms and counts numeric citations.
    return len(re.findall(r"\b[\w]+(?:[’'][\w]+)*\b", s))


def gather_evidence():
    result = {
        "date": "2026-10-07",
        "experimental_unit": "One reported empty-gripper closure-and-return sequence",
        "contact_condition": "User-reported empty; no independent contact-force reference",
        "platform_nominal_duration_s": 10,
        "magnetic_field_unit": "microtesla",
        "uncertainty": "Total physical uncertainty not quantified; sample SD is descriptive only",
        "limits": [
            "Not five independent experiments",
            "Pose and physical opening were not independently measured",
            "No matched fixed-opening time control; cause of changes not isolated",
            "No force calibration or safe-grip conclusion",
        ],
        "platforms": [],
    }
    for name, label in [("baseline", "Baseline"), ("close_1", "Close 1"),
                        ("close_2", "Close 2"), ("close_3", "Close 3"), ("return", "Return")]:
        path = RAW / (name + ".csv")
        with path.open() as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == 100 and all(int(r["valid"]) == 1 for r in rows)
        vals = [float(r["bz_uT"]) for r in rows]
        result["platforms"].append({
            "name": name, "label": label,
            "source": str(path.relative_to(ROOT)),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "n": len(rows), "bz_mean_uT": statistics.mean(vals),
            "bz_sample_sd_uT": statistics.stdev(vals),
        })
    baseline = result["platforms"][0]["bz_mean_uT"]
    for p in result["platforms"]:
        p["delta_bz_mean_uT"] = p["bz_mean_uT"] - baseline
    (HERE / "evidence.json").write_text(json.dumps(result, indent=2) + "\n")
    return result


def validate_counts():
    counts = []
    for s in CONTENT["sections"]:
        text = " ".join(s["paragraphs"])
        item = {"section": s["heading"], "maximum": s["limit"],
                "words": words(text), "conservative_words": conservative_words(text)}
        assert item["words"] <= s["limit"], item
        assert item["conservative_words"] <= s["limit"], item
        counts.append(item)
    ai = CONTENT["ai_declaration"]
    total = sum(c["words"] for c in counts) + words(ai)
    conservative = sum(c["conservative_words"] for c in counts) + conservative_words(ai)
    assert total <= 850 and conservative <= 850
    result = {"sections": counts, "ai_declaration_words": words(ai),
              "total_including_ai": total, "conservative_total_including_ai": conservative,
              "excluded": "Section headings and bibliography; no footnotes are used",
              "count_method": "Whitespace-delimited; conservative check additionally splits compounds",
              "oral_script_words": sum(words(" ".join(s["script"])) for s in ORAL["slides"]),
              "slide_count": 2}
    (HERE / "validation.json").write_text(json.dumps(result, indent=2) + "\n")
    return result


def report_markdown(counts):
    lines = ["# " + CONTENT["title"], "", CONTENT["subject"], "",
             f"{CONTENT['author']} | Student ID: {CONTENT['student_id']}", ""]
    for s in CONTENT["sections"]:
        lines += ["## " + s["heading"], ""]
        for p in s["paragraphs"]:
            lines += [p, ""]
    lines += ["## " + CONTENT["ai_heading"], "", CONTENT["ai_declaration"], "",
              "## References", ""]
    for ref in CONTENT["references"]:
        lines += [ref, ""]
    (HERE / "ATOM_Ethics_Report.md").write_text("\n".join(lines))


def report_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.9)
    section.left_margin = section.right_margin = Cm(2.0)
    section.header_distance = section.footer_distance = Cm(.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [("Title", 23, NAVY), ("Heading 1", 12, TEAL)]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(5)
    header = section.header.paragraphs[0]
    header.text = "ENGR90037  /  Individual ethical review"
    header.style = doc.styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.text = "Yifan Wu | 1618332"
    footer.paragraph_format.tab_stops.add_tab_stop(Cm(15.8))
    footer.add_run("\t")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    for run in footer.runs:
        run.font.size = Pt(9)
    doc.add_paragraph(CONTENT["title"], "Title")
    p = doc.add_paragraph(CONTENT["subject"])
    p.paragraph_format.space_after = Pt(3)
    p = doc.add_paragraph(f"{CONTENT['author']}  |  Student ID: {CONTENT['student_id']}")
    p.paragraph_format.space_after = Pt(9)
    for i, s in enumerate(CONTENT["sections"]):
        if i == 4:
            doc.add_page_break()
        doc.add_paragraph(s["heading"], "Heading 1")
        for text in s["paragraphs"]:
            doc.add_paragraph(text)
    doc.add_paragraph(CONTENT["ai_heading"], "Heading 1")
    doc.add_paragraph(CONTENT["ai_declaration"])
    doc.add_page_break()
    doc.add_paragraph("References", "Heading 1")
    for ref in CONTENT["references"]:
        p = doc.add_paragraph(ref)
        p.paragraph_format.space_after = Pt(10)
        for run in p.runs:
            run.font.size = Pt(10)
    doc.core_properties.author = CONTENT["author"]
    doc.core_properties.title = CONTENT["title"]
    doc.core_properties.subject = CONTENT["subject"]
    doc.save(HERE / "ATOM_Ethics_Report.docx")


def report_pdf():
    for name, filename in [("Arial", "Arial.ttf"), ("Arial-Bold", "Arial Bold.ttf")]:
        pdfmetrics.registerFont(TTFont(name, str(FONT_ROOT / filename)))
    body = ParagraphStyle("Body", fontName="Arial", fontSize=11, leading=13.1,
                          spaceAfter=6, alignment=TA_LEFT)
    heading = ParagraphStyle("Heading", parent=body, fontName="Arial-Bold", fontSize=12,
                             textColor=colors.HexColor("#" + TEAL), spaceBefore=9, spaceAfter=5,
                             keepWithNext=True)
    title = ParagraphStyle("Title", parent=body, fontName="Arial-Bold", fontSize=23,
                           leading=28, textColor=colors.HexColor("#" + NAVY), spaceAfter=12)
    meta = ParagraphStyle("Meta", parent=body, fontSize=10, leading=13, spaceAfter=4)
    refstyle = ParagraphStyle("Reference", parent=body, fontSize=10, leading=12.5, spaceAfter=11)
    def para(text, style=body):
        escaped = html.escape(text)
        escaped = re.sub(r"(https://[^\s)]+)", r'<link href="\1" color="#087F83">\1</link>', escaped)
        return Paragraph(escaped, style)
    story = [para(CONTENT["title"], title), para(CONTENT["subject"], meta),
             para(f"{CONTENT['author']}  |  Student ID: {CONTENT['student_id']}", meta), Spacer(1, 7)]
    for i, s in enumerate(CONTENT["sections"]):
        if i == 4:
            story.append(PageBreak())
        story.append(para(s["heading"], heading))
        story.extend(para(p) for p in s["paragraphs"])
    story += [para(CONTENT["ai_heading"], heading), para(CONTENT["ai_declaration"]), PageBreak(),
              para("References", heading)]
    story.extend(para(r, refstyle) for r in CONTENT["references"])
    def decorate(c, d):
        w, h = d.pagesize
        c.setFont("Arial", 8)
        c.setFillColor(colors.HexColor("#526775"))
        c.drawString(56.7, h - 29, "ENGR90037  /  Individual ethical review")
        c.drawString(56.7, 27, "Yifan Wu | 1618332")
        c.drawRightString(w - 56.7, 27, str(d.page))
    pdf = SimpleDocTemplate(str(HERE / "ATOM_Ethics_Report.pdf"),
                            pagesize=(595.276, 841.89),
                            leftMargin=56.7, rightMargin=56.7, topMargin=53.85,
                            bottomMargin=53.85, title=CONTENT["title"], author=CONTENT["author"])
    pdf.build(story, onFirstPage=decorate, onLaterPages=decorate)


def chart_svg(evidence):
    # Means relative to first platform; no error bars suggesting total uncertainty.
    xs = [100, 225, 350, 475, 600]
    y = lambda v: 70 + (-v) / 7 * 225
    svg = ['<svg viewBox="0 0 700 370" role="img" aria-labelledby="chart-title chart-desc">',
           '<title id="chart-title">Mean change in Bz during one empty-gripper sequence</title>',
           '<desc id="chart-desc">Magnetic field changes without reported object contact. '
           'Values are descriptive platform means in microtesla, not force.</desc>']
    for tick in [0, -2, -4, -6]:
        yy = y(tick)
        svg += [f'<line x1="74" y1="{yy}" x2="635" y2="{yy}" stroke="#D8E1E5"/>',
                f'<text x="58" y="{yy+6}" text-anchor="end" class="axis">{tick}</text>']
    svg += ['<text x="75" y="35" class="axis-title">ΔBz (µT), relative to baseline</text>']
    points = " ".join(f"{xx},{y(p['delta_bz_mean_uT']):.3f}" for xx, p in zip(xs, evidence["platforms"]))
    svg += [f'<polyline points="{points}" fill="none" stroke="#087F83" stroke-width="4" stroke-linejoin="round"/>']
    for xx, p in zip(xs, evidence["platforms"]):
        value = p["delta_bz_mean_uT"]
        yy = y(value)
        svg += [f'<circle cx="{xx}" cy="{yy:.3f}" r="7" fill="#087F83" stroke="white" stroke-width="3"/>',
                f'<text x="{xx}" y="{yy-19:.3f}" text-anchor="middle" class="data-value">{value:.2f}</text>',
                f'<text x="{xx}" y="332" text-anchor="middle" class="axis">{p["label"]}</text>']
    svg += ['</svg>']
    return "\n".join(svg)


def html_deck(evidence):
    template = (HERE / "slides.template.html").read_text()
    notes = []
    for i, s in enumerate(ORAL["slides"]):
        notes.append({"title": s["title"], "time": s["target_time"],
                      "paragraphs": s["script"], "sources": CONTENT["references"]})
    deck = template.replace("@@CHART@@", chart_svg(evidence))
    deck = deck.replace("@@NOTES_JSON@@", json.dumps(notes, ensure_ascii=False).replace("</", "<\\/"))
    deck = deck.replace("@@AI_DECLARATION@@", html.escape(CONTENT["ai_declaration"]))
    assert "@@" not in deck
    (HERE / "ATOM_Ethics_Slides.html").write_text(deck)


def speaker_notes():
    lines = ["# ATOM Ethics — Speaker Notes", "", "Yifan Wu | 1618332", "",
             "简单英文版。标题不用读，只读下面的英文句子。每句一行，段落之间可稍作停顿。", "",
             "Practice target: about 4:30; maximum: 5:00. Two slides. Check the time at your own speaking speed.", "",
             "[中英对照和词义](ATOM_Ethics_Speaker_Notes_Bilingual.md)", ""]
    for i, s in enumerate(ORAL["slides"]):
        lines += [f"## Slide {i+1}: {s['title']}", "", f"Target time: {s['target_time']}", ""]
        for p in s["script"]:
            lines += [sentence + "  " for sentence in re.split(r"(?<=[.!?])\s+", p)] + [""]
    lines += ["## AI use", "", CONTENT["ai_declaration"], ""]
    (HERE / "ATOM_Ethics_Speaker_Notes.md").write_text("\n".join(lines))


def bilingual_speaker_notes():
    count = sum(words(" ".join(s["script"])) for s in ORAL["slides"])
    lines = ["# ATOM Ethics：简单英文讲稿（中英对照）", "", "Yifan Wu · 1618332", "",
             f"英文共 {count} 词。只读英文，中文帮助理解。练习目标约 4 分 30 秒，以自己实际计时为准；最多 5 分钟。", "",
             "标题、时间提示和最后的词表不用念。先读一句英文，再确认自己知道它的中文意思。两个守则的英文说明是口头概括，并非原文引用。", ""]
    for i, s in enumerate(ORAL["slides"]):
        assert len(s["script"]) == len(s["script_zh"])
        lines += [f"## 第 {i+1} 页", "", f"练习时间：{s['target_time']}", ""]
        for en, zh in zip(s["script"], s["script_zh"]):
            lines += ["> " + sentence + "  " for sentence in re.split(r"(?<=[.!?])\s+", en)]
            lines += ["", zh, ""]
    lines += ["## 简单说法与课堂术语", "",
              "演讲时可以用下面的简单说法解释原则；老师追问名称时，再对应到报告中的术语。", "",
              "| 讲稿里的简单说法 | 报告里的原则 | 中文意思 |",
              "| --- | --- | --- |",
              "| Help people | Beneficence | 为别人带来益处 |",
              "| Avoid harm | Non-maleficence | 避免伤害 |",
              "| Be fair | Justice | 公平分配收益和负担 |", "",
              "## 需要认得的词", "",
              "| 英文 | 中文 |", "| --- | --- |",
              "| ethics | 伦理：应该怎样做，以及为什么 |",
              "| lab | 实验室，laboratory 的简称 |",
              "| gripper | 夹爪 |", "| sensor reading | 传感器读数 |",
              "| magnetic field | 磁场 |", "| force | 力 |",
              "| risk | 风险 |", "| avoid harm | 避免伤害 |",
              "| fair | 公平的 |", "| code / rule | 这里指职业守则／条款 |",
              "| environment | 环境 |", "| hardware emergency stop | 硬件急停 |", "",
              "## 只记住这条逻辑也能把第二页讲顺", "",
              "空夹爪的读数也会变化 → 磁场不是直接测出的力 → 还需验证 → 保留力控制限制 → 再检查停止和样品保持。", "",
              "讲到真实项目状态时，保留原稿中的 may、need 和 not yet：可能的影响不等于已经发生，需要完成的测试也不能说成已经完成。", ""]
    (HERE / "ATOM_Ethics_Speaker_Notes_Bilingual.md").write_text("\n".join(lines))


def main():
    counts = validate_counts()
    evidence = gather_evidence()
    report_markdown(counts)
    report_docx()
    report_pdf()
    html_deck(evidence)
    speaker_notes()
    bilingual_speaker_notes()
    print(json.dumps({"report_words_including_ai": counts["total_including_ai"],
                      "conservative_report_words": counts["conservative_total_including_ai"],
                      "slides": 2, "script_words": counts["oral_script_words"],
                      "raw_platforms": len(evidence["platforms"])}, indent=2))


if __name__ == "__main__":
    main()
